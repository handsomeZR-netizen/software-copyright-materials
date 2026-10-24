import json
from pathlib import Path
import tempfile
import unittest
import docx
from agreement import load, build, check_id

class AgreementTests(unittest.TestCase):
    def config(self, folder, **extra):
        cfg = {'software_name': '示例软件', 'development_date': '2026 年 9 月 27 日', 'signing_date': '2026 年 6 月 15 日',
               'purpose': '本软件用于示例数据管理。', 'output': 'out/协议.docx',
               'parties': [{'name': '甲某', 'duties': '服务端开发'}, {'name': '乙某', 'duties': '前端开发'}]}
        cfg.update(extra)
        path = Path(folder)/'agreement.json'
        path.write_text(json.dumps(cfg, ensure_ascii=False), encoding='utf-8')
        return path
    def text(self, path):
        d = docx.Document(path)
        return '\n'.join([p.text for p in d.paragraphs]+[c.text for c in d.tables[0]._cells]+
                         [p.text for s in d.sections for p in s.header.paragraphs])
    def test_two_parties_filled_without_placeholders(self):
        with tempfile.TemporaryDirectory() as folder:
            text = self.text(build(load(self.config(folder))))
            self.assertNotIn('XXX', text)
            self.assertNotIn('合作方三', text)
            self.assertNotIn('具体分工可由', text)
            self.assertIn('本软件著作权由甲某、乙某共同享有', text)
            self.assertIn('登记申请经办人：甲某', text)
            self.assertIn('一式两份', text)
            self.assertIn('协议签署日期：2026 年 6 月 15 日', text)
    def test_three_parties_keep_blank_fourth_cell(self):
        with tempfile.TemporaryDirectory() as folder:
            parties = [{'name': n, 'duties': '开发'} for n in ['甲某', '乙某', '丙某']]
            text = self.text(build(load(self.config(folder, parties=parties))))
            self.assertIn('一式三份', text)
            self.assertNotIn('合作方四', text)
            self.assertNotIn('XXX', text)
    def test_signing_date_must_precede_completion(self):
        with tempfile.TemporaryDirectory() as folder:
            for date in ['2026 年 9 月 27 日', '2026 年 10 月 1 日']:
                with self.subTest(date=date), self.assertRaises(ValueError):
                    load(self.config(folder, signing_date=date))
    def test_id_checksum(self):
        self.assertTrue(check_id('11010519491231002X'))
        self.assertFalse(check_id('110105194912310021'))

if __name__ == '__main__':
    unittest.main()
