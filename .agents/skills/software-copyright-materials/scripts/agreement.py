"""Fill the cooperation-agreement DOCX template from a JSON config.

python agreement.py --config agreement.json [--pdf]

Config keys (paths relative to the config file):
  software_name, software_version (default V1.0), development_date ("2026 年 9 月 27 日"),
  signing_date (must be earlier than development_date; empty leaves a blank to fill by hand),
  purpose (one sentence ending with 。), parties [{name, id_number, duties}] (2-4, signing order),
  handler (default: first party), output (.docx), template (optional).
"""
import argparse, json, pathlib, re, subprocess, sys
import docx

KIT = pathlib.Path(__file__).resolve().parent.parent
TEMPLATE = KIT / 'assets/agreement-template/合作开发协议模板.docx'
ORDINALS = '一二三四'
BLANK_DATE = '        年      月      日'


def parse_date(text):
    m = re.fullmatch(r'\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*', text or '')
    if not m:
        raise ValueError('Date must look like "2026 年 9 月 27 日": '+repr(text))
    return tuple(int(x) for x in m.groups())


def check_id(number):
    weights = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
    if not re.fullmatch(r'\d{17}[\dXx]', number):
        return False
    return '10X98765432'[sum(int(a)*b for a, b in zip(number, weights)) % 11] == number[-1].upper()


def load(path):
    cfg = json.loads(path.read_text(encoding='utf-8'))
    for key in ['software_name', 'development_date', 'purpose', 'parties', 'output']:
        if not cfg.get(key):
            raise ValueError('Missing '+key)
    parties = cfg['parties']
    if not 2 <= len(parties) <= 4:
        raise ValueError('Template supports 2-4 parties')
    for p in parties:
        if not p.get('name') or not p.get('duties'):
            raise ValueError('Each party needs name and duties')
        if p.get('id_number') and not check_id(p['id_number']):
            raise ValueError('ID number checksum failed: '+p['name'])
    developed = parse_date(cfg['development_date'])
    if cfg.get('signing_date') and parse_date(cfg['signing_date']) >= developed:
        raise ValueError('signing_date must be earlier than development_date')
    cfg.setdefault('software_version', 'V1.0')
    cfg.setdefault('handler', parties[0]['name'])
    cfg['_base'] = path.parent
    return cfg


def fill(run, text):
    run.text = text
    run.font.color.rgb = None


def drop(element):
    element._element.getparent().remove(element._element)


def placeholders(paragraph):
    return [r for r in paragraph.runs if r.text == 'XXX']


def set_version(paragraph, version):
    for run in paragraph.runs:
        if 'V1.0' in run.text:
            run.text = run.text.replace('V1.0', version)


def build(cfg):
    template = cfg['_base'] / cfg['template'] if cfg.get('template') else TEMPLATE
    d = docx.Document(template)
    name, version, parties = cfg['software_name'], cfg['software_version'], cfg['parties']
    signing = cfg.get('signing_date') or BLANK_DATE
    n = len(parties)
    for para in list(d.paragraphs):
        text, slots = para.text, placeholders(para)
        set_version(para, version)
        if text.startswith(('本协议由以下合作方', '软件全称及版本号：')):
            fill(slots[0], name)
        elif text.startswith('开发完成日期：'):
            fill(slots[0], cfg['development_date'])
        elif text.startswith('协议签署日期：'):
            fill(slots[0], signing)
        elif text.startswith('登记申请经办人：') or text.startswith('协议各方同意由XXX'):
            fill(slots[0], cfg['handler'])
        elif text.startswith('协议各方决定共同开发'):
            fill(slots[0], name)
            fill(slots[1], cfg['purpose'])
        elif text.startswith('协议各方一致同意，本软件著作权由'):
            fill(slots[0], '、'.join(p['name'] for p in parties))
            first = [r._r for r in para.runs].index(slots[0]._r)
            for run in para.runs[first+1:]:
                if run.text in ('XXX', '、'):
                    drop(run)
        elif re.match('合作方[一二三四][：（]', text):
            index = ORDINALS.index(text[3])
            if index >= n:
                drop(para)
                continue
            party = parties[index]
            if text[4] == '：':
                fill(slots[0], party['name']+('（身份证号码：'+party['id_number']+'）' if party.get('id_number') else ''))
            else:
                fill(slots[0], party['name'])
                fill(slots[1], party['duties'].rstrip('。'))
        elif '一式四份' in text:
            for run in para.runs:
                run.text = run.text.replace('一式四份', '一式'+'一两三四'[n-1]+'份')

    table = d.tables[0]
    cells = [c for row in table.rows for c in row.cells]
    for index, cell in enumerate(cells):
        if index < n:
            for para in cell.paragraphs:
                if para.text.startswith('名称/姓名：'):
                    fill(placeholders(para)[0], parties[index]['name'])
        else:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.text = ''
    if n <= 2:
        table._tbl.remove(table.rows[1]._tr)

    for section in d.sections:
        for para in section.header.paragraphs:
            set_version(para, version)
            for run in placeholders(para):
                fill(run, name)
    props = d.core_properties
    props.title = name+' '+version+' 合作开发协议'
    props.author = props.last_modified_by = props.comments = ''

    leftover = [p.text for p in d.paragraphs if 'XXX' in p.text]
    leftover += [c.text for c in table._cells if 'XXX' in c.text]
    leftover += [p.text for s in d.sections for p in s.header.paragraphs if 'XXX' in p.text]
    if leftover:
        raise ValueError('Unfilled placeholders: '+'; '.join(leftover))
    out = cfg['_base'] / cfg['output']
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(out)
    return out


def export_pdf(path):
    """Refresh page fields and export with Microsoft Word (Windows)."""
    script = ('$w=New-Object -ComObject Word.Application;$w.Visible=$false;try{'
              '$d=$w.Documents.Open($env:AGREEMENT_DOCX);$d.Fields.Update()|Out-Null;$d.Save();'
              '$d.ExportAsFixedFormat($env:AGREEMENT_PDF,17);$d.Close()}'
              'finally{if($w.Documents.Count -eq 0){$w.Quit()}}')
    pdf = path.with_suffix('.pdf')
    env = dict(__import__('os').environ, AGREEMENT_DOCX=str(path), AGREEMENT_PDF=str(pdf))
    subprocess.run(['powershell', '-NoProfile', '-Command', script], check=True, env=env)
    return pdf


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--config', required=True)
    parser.add_argument('--pdf', action='store_true', help='also export PDF via Word')
    args = parser.parse_args()
    out = build(load(pathlib.Path(args.config).resolve()))
    print(out)
    if args.pdf:
        print(export_pdf(out))


if __name__ == '__main__':
    sys.exit(main())
